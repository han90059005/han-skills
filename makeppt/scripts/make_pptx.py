#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""make_pptx.py —— 根据 slides.json 生成一份 16:9 NWPU 风格 .pptx 汇报。

用法：
    python make_pptx.py slides.json -o 输出.pptx

slides.json 字段（除 slides[].title 外均可省略）：
{
  "title": "主标题",
  "subtitle": "副标题",
  "author": "姓名",
  "affiliation": "单位",
  "date": "日期/备注",
  "logo": "标题页 logo 图片路径",
  "slides": [
    {
      "title": "页标题",
      "bullets": ["要点1", "要点2"],
      "table": {"headers": ["列1", "列2"], "rows": [["a", "b"], ["c", "d"]]},
      "image": "图片路径",
      "caption": "图注"
    }
  ]
}

硬规则落地：文字颜色 ≤ 3 种（正文黑、标题白、蓝色只做背景/色块）；
每页右下加页码「x / y」；正文默认 18pt、标题 26pt，保证清晰可读。
"""
import argparse
import json
import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# NWPU 配色
BLUE = RGBColor(0x23, 0x68, 0xB1)       # 35,104,177
BLUE_DARK = RGBColor(0x1B, 0x50, 0x8A)
BLUE_GRAY = RGBColor(0xE5, 0xED, 0xF6)  # 229,237,246
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "微软雅黑"
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def set_run_font(run, size, bold=False, color=BLACK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    # 东亚字体：保证中文用指定字体而不是系统默认
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", FONT)


def add_rect(slide, x, y, w, h, fill):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_text(slide, x, y, w, h, text, size=18, bold=False, color=BLACK,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.15):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    first = True
    for para in text.split("\n"):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.line_spacing = line_spacing
        run = p.add_run()
        run.text = para
        set_run_font(run, size, bold, color)
    return box


def add_bullets(slide, x, y, w, h, bullets, size=18):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    first = True
    for b in bullets:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.line_spacing = 1.2
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = "\u2022  " + str(b)
        set_run_font(run, size, False, BLACK)
    return box


def add_table(slide, x, y, w, headers, rows):
    n_rows = len(rows) + 1
    n_cols = len(headers)
    gf = slide.shapes.add_table(n_rows, n_cols, x, y, w, Inches(0.4 * n_rows))
    table = gf.table
    col_w = int(w / n_cols)
    for i in range(n_cols):
        table.columns[i].width = col_w
    # 表头：蓝底白字
    for j, htxt in enumerate(headers):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = BLUE
        cell.margin_top = Pt(3)
        cell.margin_bottom = Pt(3)
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = str(htxt)
        set_run_font(run, 14, True, WHITE)
    # 表体：白 / 浅蓝交替
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if i % 2 == 1 else BLUE_GRAY
            cell.margin_top = Pt(3)
            cell.margin_bottom = Pt(3)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.text = str(val)
            set_run_font(run, 13, False, BLACK)
    return gf


def add_title_slide(prs, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, BLUE)
    add_rect(slide, 0, 0, SLIDE_W, Inches(0.14), BLUE_DARK)

    if spec.get("logo") and os.path.exists(spec["logo"]):
        slide.shapes.add_picture(spec["logo"], Inches(6.42), Inches(0.45), height=Inches(0.9))

    add_text(slide, Inches(0.8), Inches(2.0), Inches(11.73), Inches(1.4),
             spec.get("title", ""), size=40, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    if spec.get("subtitle"):
        add_text(slide, Inches(0.8), Inches(3.5), Inches(11.73), Inches(0.7),
                 spec["subtitle"], size=24, color=WHITE, align=PP_ALIGN.CENTER)

    meta = []
    if spec.get("author"):
        meta.append(spec["author"])
    if spec.get("affiliation"):
        meta.append(spec["affiliation"])
    if spec.get("date"):
        meta.append(spec["date"])
    if meta:
        add_text(slide, Inches(0.8), Inches(5.1), Inches(11.73), Inches(1.9),
                 "\n".join(meta), size=16, color=WHITE,
                 align=PP_ALIGN.CENTER, line_spacing=1.5)


def add_content_slide(prs, spec, idx, total, base_dir):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 标题栏
    add_rect(slide, 0, 0, SLIDE_W, Inches(1.0), BLUE)
    add_text(slide, Inches(0.5), Inches(0.1), Inches(12.3), Inches(0.8),
             spec.get("title", ""), size=26, bold=True, color=WHITE,
             anchor=MSO_ANCHOR.MIDDLE)

    # 页码
    add_text(slide, Inches(11.85), Inches(7.0), Inches(1.25), Inches(0.4),
             f"{idx} / {total}", size=12, color=BLACK, align=PP_ALIGN.RIGHT)

    body_top = 1.25
    body_bottom = 6.85
    body_h = body_bottom - body_top

    bullets = spec.get("bullets") or []
    table = spec.get("table")
    image = spec.get("image")
    caption = spec.get("caption")

    # 粗估各块高度（英寸）
    bullets_h = min(0.5 * len(bullets), body_h * 0.55) if bullets else 0.0
    table_h = 0.44 * (len(table["rows"]) + 1) if table else 0.0
    caption_h = 0.42 if caption else 0.0
    image_h = 0.0
    if image:
        image_h = body_h - bullets_h - table_h - caption_h - 0.25
        image_h = max(image_h, 0.8)

    y = body_top
    if bullets:
        add_bullets(slide, Inches(0.6), Inches(y), Inches(12.1), Inches(bullets_h), bullets)
        y += bullets_h + 0.15
    if table:
        add_table(slide, Inches(0.6), Inches(y), Inches(12.1), table["headers"], table["rows"])
        y += table_h + 0.15
    if image:
        img_path = image
        if not os.path.isabs(img_path):
            img_path = os.path.join(base_dir, img_path)
        if os.path.exists(img_path):
            pic = slide.shapes.add_picture(img_path, Inches(0.6), Inches(y), height=Inches(image_h))
            scale = min(Inches(12.0) / pic.width, Inches(image_h) / pic.height)
            pic.width = int(pic.width * scale)
            pic.height = int(pic.height * scale)
            pic.left = int((SLIDE_W - pic.width) / 2)
            y += image_h + 0.1
        else:
            add_text(slide, Inches(0.6), Inches(y), Inches(12.1), Inches(0.6),
                     f"[缺图] {image}", size=16, color=BLACK)
            y += 0.6
    if caption:
        add_text(slide, Inches(0.6), Inches(y), Inches(12.1), Inches(caption_h),
                 caption, size=13, color=BLACK, align=PP_ALIGN.CENTER)


def main():
    ap = argparse.ArgumentParser(description="根据 slides.json 生成 NWPU 风格 .pptx")
    ap.add_argument("spec", help="slides.json 路径")
    ap.add_argument("-o", "--out", required=True, help="输出 .pptx 路径")
    args = ap.parse_args()

    with open(args.spec, "r", encoding="utf-8") as f:
        spec = json.load(f)

    base_dir = os.path.dirname(os.path.abspath(args.spec))

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    add_title_slide(prs, spec)

    slides = spec.get("slides", [])
    for i, s in enumerate(slides, start=1):
        add_content_slide(prs, s, i, len(slides), base_dir)

    out = args.out
    out_dir = os.path.dirname(os.path.abspath(out))
    if out_dir and not os.path.isdir(out_dir):
        os.makedirs(out_dir, exist_ok=True)
    prs.save(out)
    print(f"[makeppt] pptx 已生成：{out}")


if __name__ == "__main__":
    main()
